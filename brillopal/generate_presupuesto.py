#!/usr/bin/env python3
"""Generate BrillOPal_Presupuesto_Tipo.docx"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import date

BLUE = RGBColor(0x1F, 0x38, 0x64)
GOLD = RGBColor(0xF4, 0xB9, 0x42)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0xF2, 0xF2, 0xF2)

def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    if level == 1:
        run = p.runs[0]
        run.font.color.rgb = WHITE
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), '1F3864')
        p._p.get_or_add_pPr().append(shd)
    return p

def add_para(doc, text, bold=False, italic=False, color=None, size=10, align=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    p.paragraph_format.space_after = Pt(3)
    return p

def add_table(doc, headers, rows, col_widths=None, header_bg='1F3864', alt_bg='D6E4F0'):
    table = doc.add_table(rows=1+len(rows), cols=len(headers))
    table.style = 'Table Grid'
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(9)
        cell.paragraphs[0].runs[0].font.color.rgb = WHITE
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_bg(cell, header_bg)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i+1].cells[c_i]
            cell.text = str(val) if val is not None else ""
            cell.paragraphs[0].runs[0].font.size = Pt(9)
            if r_i % 2 == 1:
                set_cell_bg(cell, alt_bg)
    if col_widths:
        for i, w in enumerate(col_widths):
            for cell in table.columns[i].cells:
                cell.width = Cm(w)
    return table

def separator(doc):
    p = doc.add_paragraph("─" * 80)
    p.runs[0].font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    p.runs[0].font.size = Pt(7)

def build_header_block(doc, tipo):
    """Builds the company + client header for a budget"""
    tbl = doc.add_table(rows=2, cols=2)
    tbl.style = 'Table Grid'

    # Company cell
    c = tbl.rows[0].cells[0]
    set_cell_bg(c, '1F3864')
    c.paragraphs[0].add_run("BRILLOPAL").bold = True
    c.paragraphs[0].runs[0].font.color.rgb = WHITE
    c.paragraphs[0].runs[0].font.size = Pt(16)
    p2 = c.add_paragraph("Servicios de Limpieza y Mantenimiento")
    p2.runs[0].font.color.rgb = WHITE
    p2.runs[0].font.size = Pt(9)
    p3 = c.add_paragraph("Palencia y Provincia | Tel.: 9XX XXX XXX\ninfo@brillopal.es | www.brillopal.es\nNIF: XXXXXXXXX")
    p3.runs[0].font.color.rgb = RGBColor(0xCC, 0xDD, 0xFF)
    p3.runs[0].font.size = Pt(8)

    # Budget ref cell
    c2 = tbl.rows[0].cells[1]
    set_cell_bg(c2, 'F4B942')
    c2.paragraphs[0].add_run(f"OFERTA / PRESUPUESTO {tipo}").bold = True
    c2.paragraphs[0].runs[0].font.color.rgb = BLUE
    c2.paragraphs[0].runs[0].font.size = Pt(11)
    ref_info = [
        ("Nº Presupuesto:", "PRES-2026-___"),
        ("Fecha emisión:", date.today().strftime("%d/%m/%Y")),
        ("Válido hasta:", "____/____/______"),
        ("Comercial:", "________________________"),
    ]
    for label, val in ref_info:
        pr = c2.add_paragraph()
        r_lab = pr.add_run(f"{label} ")
        r_lab.bold = True
        r_lab.font.size = Pt(9)
        r_lab.font.color.rgb = BLUE
        r_val = pr.add_run(val)
        r_val.font.size = Pt(9)
        r_val.font.color.rgb = BLUE

    # Client data cell
    c3 = tbl.rows[1].cells[0]
    c3.merge(tbl.rows[1].cells[1])
    set_cell_bg(c3, 'D6E4F0')
    c3.paragraphs[0].add_run("DATOS DEL CLIENTE").bold = True
    c3.paragraphs[0].runs[0].font.color.rgb = BLUE
    c3.paragraphs[0].runs[0].font.size = Pt(10)
    fields = [
        ("Empresa/Nombre:", "____________________________________________"),
        ("NIF/CIF:",         "____________________________________________"),
        ("Dirección:",       "____________________________________________"),
        ("Teléfono:",        "____________________  Email: _________________________"),
        ("Contacto:",        "____________________________________________"),
    ]
    for label, blank in fields:
        pf = c3.add_paragraph()
        r_l = pf.add_run(f"{label}  ")
        r_l.bold = True
        r_l.font.size = Pt(9)
        r_l.font.color.rgb = BLUE
        r_b = pf.add_run(blank)
        r_b.font.size = Pt(9)

    for col in tbl.columns:
        col.cells[0].width = Cm(8.5)


# ──────────────────────────────────────────────────────────────────────────────
def build_b2b(doc):
    add_heading(doc, "OFERTA DE SERVICIOS B2B – SUBCONTRATACIÓN EMPRESAS")

    add_para(doc, "Estimado cliente:", bold=True)
    add_para(doc,
        "En respuesta a su solicitud, nos complace presentarle la siguiente propuesta de "
        "servicios de limpieza y mantenimiento para sus instalaciones. BRILLOPAL se "
        "compromete a gestionar íntegramente el servicio con personal propio cualificado, "
        "cumplimiento total de la normativa PRL, coordinación CAE disponible en 48h y "
        "garantía de servicio con sustituciones en <24h.")

    doc.add_paragraph()
    build_header_block(doc, "B2B")
    doc.add_paragraph()

    # Service detail table
    add_para(doc, "DETALLE DEL SERVICIO PROPUESTO", bold=True, color=BLUE)
    add_table(doc,
        ["Servicio", "Descripción detallada", "Frecuencia", "H/sem", "H/mes", "€/hora\n(s/IVA)", "€/mes\n(s/IVA)"],
        [
            ("Limpieza general", "Suelos, superficies, mobiliario, cristales interiores, papeleras, aseos", "5 días/sem", "10,0", "40,0", "20,41", "816,40"),
            ("Limpieza profunda", "Zonas de difícil acceso, maquinaria externa, vestuarios industriales", "1 día/sem", "3,0", "12,0", "20,41", "244,92"),
            ("Servicio adicional 1", "____________________________", "______", "___", "___", "______", "______"),
            ("Servicio adicional 2", "____________________________", "______", "___", "___", "______", "______"),
        ],
        [3, 5.5, 2.5, 1.2, 1.2, 1.8, 1.8])
    doc.add_paragraph()

    # Totals
    totals_tbl = doc.add_table(rows=4, cols=4)
    totals_tbl.style = 'Table Grid'
    totals_data = [
        ("Total horas/mes contratadas:", "52,0 h/mes", "Precio/hora:", "20,41 €/h"),
        ("BASE IMPONIBLE MENSUAL:", "1.061,32 €", "", ""),
        ("IVA (21%):", "222,88 €", "", ""),
        ("TOTAL MENSUAL (c/IVA):", "1.284,20 €", "", ""),
    ]
    for r_i, (l1, v1, l2, v2) in enumerate(totals_data):
        for c_i, val in enumerate([l1, v1, l2, v2]):
            cell = totals_tbl.rows[r_i].cells[c_i]
            cell.text = val
            if val and (val.startswith("BASE") or val.startswith("TOTAL MENSUAL")):
                set_cell_bg(cell, '1F3864')
                cell.paragraphs[0].runs[0].font.color.rgb = WHITE
                cell.paragraphs[0].runs[0].bold = True
            elif val and val.endswith("€") and not val.startswith("20"):
                set_cell_bg(cell, 'F4B942')
                cell.paragraphs[0].runs[0].bold = True
            cell.paragraphs[0].runs[0].font.size = Pt(9)
    doc.add_paragraph()

    # Conditions
    add_para(doc, "CONDICIONES GENERALES", bold=True, color=BLUE)
    conditions = [
        ("Duración del contrato:", "12 meses, prorrogable automáticamente por períodos iguales."),
        ("Revisión de precio:", "Anual, s/IPC del año anterior + variación del convenio colectivo de limpieza."),
        ("Preaviso de cancelación:", "60 días naturales. Penalización del 10% de la facturación pendiente hasta fin de contrato."),
        ("Subrogación:", "En cumplimiento del Convenio Colectivo de Limpieza de Palencia, si existe plantilla adscrita al servicio, BRILLOPAL procederá a su subrogación previa comunicación y análisis de costes."),
        ("Gestión de residuos:", "Los residuos generados durante el servicio serán gestionados conforme a la normativa aplicable."),
        ("PRL y CAE:", "BRILLOPAL aporta documentación CAE completa (evaluación de riesgos, seguros, formación) en un plazo máximo de 48h desde firma del contrato."),
        ("Productos utilizados:", "Productos homologados, biodegradables y con fichas de seguridad disponibles. El cliente podrá especificar restricciones de productos por escrito."),
        ("Responsabilidad civil:", "BRILLOPAL dispone de seguro de RC por importe mínimo de 300.000 €. Póliza disponible para el cliente bajo petición."),
        ("Facturación:", "Mensual, por servicios prestados. Vencimiento a 30 días desde fecha de factura. Domiciliación bancaria preferente."),
        ("Resolución de incidencias:", "Comunicación de incidencias en <2h mediante canal de contacto directo. Resolución garantizada en <24h."),
    ]
    for label, text in conditions:
        p = doc.add_paragraph()
        r_l = p.add_run(f"{label} ")
        r_l.bold = True
        r_l.font.size = Pt(9)
        r_l.font.color.rgb = BLUE
        r_t = p.add_run(text)
        r_t.font.size = Pt(9)
        p.paragraph_format.space_after = Pt(2)

    doc.add_paragraph()

    # Signature block
    add_para(doc, "ACEPTACIÓN DE OFERTA", bold=True, color=BLUE)
    sig_tbl = doc.add_table(rows=2, cols=2)
    sig_tbl.style = 'Table Grid'
    for r_i, row in enumerate([
        ("Por BRILLOPAL:", "Por la empresa cliente:"),
        ("\n\n\nFirma y sello\nNombre: _____________________\nFecha: ______________________",
         "\n\n\nFirma y sello\nNombre: _____________________\nDNI/NIF: ____________________\nFecha: ______________________"),
    ]):
        for c_i, val in enumerate(row):
            cell = sig_tbl.rows[r_i].cells[c_i]
            cell.text = val
            cell.paragraphs[0].runs[0].font.size = Pt(9)
            if r_i == 0:
                cell.paragraphs[0].runs[0].bold = True
                set_cell_bg(cell, 'D6E4F0')


# ──────────────────────────────────────────────────────────────────────────────
def build_b2c(doc):
    doc.add_page_break()
    add_heading(doc, "PRESUPUESTO B2C – LIMPIEZA DE HOGAR Y PEQUEÑAS OFICINAS")

    add_para(doc, "Estimado/a cliente/a:", bold=True)
    add_para(doc,
        "Le presentamos a continuación nuestra propuesta de servicio de limpieza doméstica. "
        "BRILLOPAL ofrece un servicio personalizado, con personal fijo asignado a su hogar, "
        "productos de alta calidad y total discreción y confianza.")

    doc.add_paragraph()
    build_header_block(doc, "B2C")
    doc.add_paragraph()

    add_para(doc, "DETALLE DEL SERVICIO", bold=True, color=BLUE)
    add_table(doc,
        ["Servicio", "Descripción", "H/visita", "Visitas/mes", "€/hora", "€/mes"],
        [
            ("Limpieza estándar", "Barrido+fregado, cocina, baños, habitaciones, polvo", "3,0", "2", "22,50", "135,00"),
            ("Limpieza profunda", "Todo lo anterior + interior electrodomésticos, armarios, ventanas", "5,0", "1 (primer mes)", "22,50", "112,50"),
            ("Plancha", "Plancha de ropa por bolsa/cesta", "—", "—", "20,00/h", "________"),
            ("Otro:", "________________________", "___", "___", "______", "______"),
        ],
        [3.5, 5, 1.5, 2, 1.5, 1.5])
    doc.add_paragraph()

    # Simple totals
    totals = [
        ("Horas totales/mes:", "6,0 h"),
        ("BASE IMPONIBLE (s/IVA):", "135,00 €"),
        ("IVA (21%):", "28,35 €"),
        ("TOTAL MENSUAL:", "163,35 €"),
    ]
    for label, val in totals:
        p = doc.add_paragraph()
        r_l = p.add_run(f"{label}  ")
        r_l.bold = True
        r_l.font.size = Pt(10)
        r_l.font.color.rgb = BLUE
        r_v = p.add_run(val)
        r_v.bold = ("TOTAL" in label)
        r_v.font.size = Pt(10)
        if "TOTAL MENSUAL" in label:
            r_v.font.color.rgb = GOLD

    doc.add_paragraph()

    add_para(doc, "CONDICIONES", bold=True, color=BLUE)
    simple_conds = [
        "Precio incluye desplazamiento dentro de Palencia capital (consultar fuera de capital).",
        "Los productos de limpieza corren a cargo de BRILLOPAL salvo indicación en contrario.",
        "Horario flexible: de lunes a sábado, franja a acordar con el cliente.",
        "Mismo/a profesional asignado/a en cada visita, salvo enfermedad o vacaciones.",
        "Sustitución garantizada en caso de ausencia del profesional habitual.",
        "Primera limpieza: se recomienda limpieza profunda inicial (1ª visita). Presupuesto adicional si se requiere.",
        "Pago: mensual, al finalizar el mes de servicio. Transferencia o domiciliación bancaria.",
        "Cancelación de visita: avisar con 24h de antelación para evitar cargo del 50% de la sesión.",
        "Presupuesto válido por 30 días desde la fecha de emisión.",
    ]
    for c in simple_conds:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(c).font.size = Pt(9)
        p.paragraph_format.space_after = Pt(2)

    doc.add_paragraph()
    separator(doc)
    add_para(doc, "¡Gracias por confiar en BRILLOPAL! Para confirmar el servicio o resolver cualquier duda, "
             "contacte con nosotros sin compromiso.", italic=True, size=9, color=BLUE)


def main():
    doc = Document()

    for section in doc.sections:
        section.top_margin    = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin   = Cm(2.2)
        section.right_margin  = Cm(2.2)

    build_b2b(doc)
    build_b2c(doc)

    doc.save("/home/user/cv-linkedin/brillopal/BrillOPal_Presupuesto_Tipo.docx")
    print("Saved: BrillOPal_Presupuesto_Tipo.docx")

if __name__ == "__main__":
    main()
